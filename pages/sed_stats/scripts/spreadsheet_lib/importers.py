"""Parsers that turn uploaded spreadsheet files into a list of row dicts."""
import csv
import io

SUPPORTED_EXTENSIONS = (".csv", ".xls", ".xlsx")


def _rows_to_dicts(rows):
    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    if not rows:
        return []

    headers = [
        str(header) if header not in (None, "") else f"column_{i + 1}"
        for i, header in enumerate(rows[0])
    ]
    result = []
    for row in rows[1:]:
        padded = list(row) + [""] * (len(headers) - len(row))
        result.append({
            header: "" if value is None else value
            for header, value in zip(headers, padded)
        })
    return result


def parse_csv(text: str):
    return _rows_to_dicts(list(csv.reader(io.StringIO(text))))


def parse_xlsx(content: bytes):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    return _rows_to_dicts([list(row) for row in sheet.iter_rows(values_only=True)])


def parse_xls(content: bytes):
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    return _rows_to_dicts([sheet.row_values(i) for i in range(sheet.nrows)])


def parse_file(filename: str, content):
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv(content)
    if lower.endswith(".xlsx"):
        return parse_xlsx(content)
    if lower.endswith(".xls"):
        return parse_xls(content)
    raise ValueError(
        f"Unsupported file type: {filename} (supported: {', '.join(SUPPORTED_EXTENSIONS)})"
    )

def rows_to_csv(rows) -> str:
    """
    Turn row dicts back into CSV text.

    The inverse of `parse_csv`, for handing the grid's contents to something
    that wants a file. Column order follows the first row, with any key a later
    row introduces appended, so a ragged set of rows still round-trips.
    """
    rows = list(rows)
    if not rows:
        return ""

    headers = list(rows[0].keys())
    seen = set(headers)
    for row in rows[1:]:
        for key in row:
            if key not in seen:
                seen.add(key)
                headers.append(key)

    buffer = io.StringIO()
    # Excel and Sheets both expect CRLF, and csv defaults to it; asking for it
    # explicitly keeps the output the same wherever this runs.
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if row.get(h) is None else row.get(h, "")
                         for h in headers])
    return buffer.getvalue()
